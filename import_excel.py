"""
import_excel.py — One-time import of the Excel workbook into SQLite.

Usage:
    python import_excel.py                       # uses Sample_data.xlsx
    python import_excel.py path/to/workbook.xlsx  # custom path
"""

import sys
import os
import openpyxl
from database import init_db, get_db, DB_PATH


def import_workbook(path: str):
    if os.path.exists(DB_PATH):
        print(f"Database already exists at {DB_PATH}")
        resp = input("Overwrite? (y/N): ").strip().lower()
        if resp != "y":
            print("Aborted.")
            return
        os.remove(DB_PATH)

    init_db()
    wb = openpyxl.load_workbook(path, data_only=True)
    conn = get_db()

    # ── Ranked Pairs ──────────────────────────────────────────────────
    ws = wb["Ranked Pairs"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))  # skip header
    for i, row in enumerate(rows):
        rank, professor, lab_group, primary_area, key_dirs, optionality, \
            prof_value, sem_suit, score, why = row[:10]
        conn.execute(
            """INSERT INTO research_groups
               (rank, professor, college, lab_group,
                primary_research_area, key_research_directions,
                professor_research_value,
                overall_score, why_it_ranks_here)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (
                int(rank) if rank else i + 1,
                str(professor or "").strip(),
                "IIIT Hyderabad",
                str(lab_group or "").strip(),
                str(primary_area or "").strip(),
                str(key_dirs or "").strip(),
                str(prof_value or "").strip(),
                round(float(score), 2) if score else None,
                str(why or "").strip(),
            ),
        )
    print(f"Imported {len(rows)} research groups.")

    # ── Field Optionality ─────────────────────────────────────────────
    ws = wb["Field Optionality"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for row in rows:
        field, precog, mll, cvit, ltrc, rrc, csg, serc, comments = row[:9]
        conn.execute(
            """INSERT INTO field_optionality
               (field, precog, mll, cvit, ltrc, rrc, csg, serc, comments)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            tuple(str(v or "").strip() for v in
                  (field, precog, mll, cvit, ltrc, rrc, csg, serc, comments)),
        )
    print(f"Imported {len(rows)} field-optionality rows.")

    # ── Lab Overview ──────────────────────────────────────────────────
    ws = wb["Lab Overview"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for row in rows:
        overall_rank, lab, identity, fields, opt, tradeoff = row[:6]
        conn.execute(
            """INSERT INTO lab_overview
               (overall_rank, lab_group, core_identity, main_fields,
                optionality, key_tradeoff)
               VALUES (?,?,?,?,?,?)""",
            (
                int(overall_rank) if overall_rank else None,
                str(lab or "").strip(),
                str(identity or "").strip(),
                str(fields or "").strip(),
                str(opt or "").strip(),
                str(tradeoff or "").strip(),
            ),
        )
    print(f"Imported {len(rows)} lab-overview rows.")

    # ── How To Decide ─────────────────────────────────────────────────
    ws = wb["How To Decide"]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    for row in rows:
        principle = row[0]
        details = row[1] if len(row) > 1 else None
        if not principle:
            continue  # skip blank rows
        conn.execute(
            "INSERT INTO how_to_decide (principle, details) VALUES (?,?)",
            (str(principle).strip(), str(details).strip() if details else None),
        )
    print(f"Imported {sum(1 for r in rows if r[0])} how-to-decide entries.")

    conn.commit()
    conn.close()
    print(f"\nDatabase created at: {DB_PATH}")


if __name__ == "__main__":
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "Sample_data.xlsx"
    )
    if not os.path.exists(xlsx):
        print(f"File not found: {xlsx}")
        sys.exit(1)
    import_workbook(xlsx)
